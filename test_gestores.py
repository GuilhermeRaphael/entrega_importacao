"""Verificações com arquivos temporários e e-mails simulados."""
import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl
import core
import gestores

class GestoresTests(unittest.TestCase):
    def vinculo(self, produto, padrao="SIM", ativo="SIM"):
        return {"PRODUTO": produto, "TIPO": "GESTOR", "PADRÃO": padrao, "ATIVO": ativo}

    def pessoa(self, nome="Ana", produtos=("AUTO", "VIDA"), email="ana@example.test"):
        return {"PESSOA": nome, "EMAIL": email, "PRODUTOS": [self.vinculo(p) for p in produtos]}

    def test_migracao_preserva_todos_vinculos_e_id(self):
        config = {"gestores": [{"PESSOA":"Ana", "EMAIL":"a@test", **self.vinculo("AUTO")},
                              {"PESSOA":"Ana", "EMAIL":"a@test", **self.vinculo("VIDA", "NÃO", "NÃO")}],
                  "vendedores":[{"nome":"V1"}]}
        gestores.migrar_cadastros_gestores(config)
        self.assertEqual(len(config["gestores"]), 1)
        self.assertEqual(len(config["gestores"][0]["PRODUTOS"]), 2)
        antes = copy.deepcopy(config)
        gestores.migrar_cadastros_gestores(config)
        self.assertEqual(config, antes)
        df = gestores.obter_configuracoes_gestores(config)
        ativos, ausentes = gestores.destinatarios_por_produto(["AUTO", "VIDA"], df)
        self.assertEqual(list(ativos["PRODUTO"]), ["AUTO"])
        self.assertEqual(ausentes, ["VIDA"])

    def test_cadastro_edicao_exclusao_persistencia(self):
        config = {"matrizes":[{"id":"M1"}], "vendedores":[{"nome":"V1"}]}
        gestores.salvar_cadastro_gestor(config, self.pessoa())
        identidade = config["gestores"][0]["ID"]
        gestores.salvar_cadastro_gestor(config, self.pessoa("Ana alterada", ("AUTO", "PET")), 0)
        self.assertEqual(config["gestores"][0]["ID"], identidade)
        with tempfile.TemporaryDirectory() as pasta:
            caminho = str(Path(pasta) / "config.json")
            core.save_config(config, caminho)
            self.assertEqual(core.load_config(caminho), config)
        with self.assertRaises(ValueError):
            gestores.salvar_cadastro_gestor(config, self.pessoa("Ana alterada"))
        duplicado = self.pessoa(produtos=("AUTO", "AUTO"))
        with self.assertRaises(ValueError):
            gestores.salvar_cadastro_gestor(config, duplicado, 0)
        gestores.excluir_cadastro_gestor(config, 0)
        self.assertEqual(config["gestores"], [])
        self.assertEqual(config["matrizes"], [{"id":"M1"}])

    def test_relatorio_unico_por_pessoa_e_reuso_por_produtos(self):
        config = {}
        for cadastro in [self.pessoa(), self.pessoa("Bia", email="bia@test"),
                         self.pessoa("Carlos", ("AUTO",), email=""), self.pessoa("Dora", ("AUTO",), email="dora@test")]:
            gestores.salvar_cadastro_gestor(config, cadastro)
        selecionados = gestores.obter_configuracoes_gestores(config).to_dict("records")
        selecionados += [dict(selecionados[0])]
        leads = [{"n":1, "nome":"Cliente 1", "responsavel":"V1", "produto":"AUTO", "empresa":"E1", "telefone":"111", "cpf":"123"},
                 {"n":2, "nome":"Cliente 2", "responsavel":"V2", "produto":"AUTO", "empresa":"E2"},
                 {"n":3, "nome":"Cliente 3", "responsavel":"V3", "produto":"VIDA", "empresa":"E3", "bonif":"BONIFICACAO"},
                 {"n":4, "nome":"Cliente 4", "responsavel":"V4", "produto":"PET", "empresa":"E4"}]
        with tempfile.TemporaryDirectory() as pasta:
            relatorios = gestores.gerar_relatorios_gerais(leads, selecionados, pasta)
            self.assertEqual(len(relatorios), 4)
            self.assertEqual(len(list(Path(pasta).rglob("*.xlsx"))), 2)
            ana, bia, carlos, dora = [relatorios[r["ID"]] for r in config["gestores"]]
            self.assertEqual(ana["caminho"], bia["caminho"])
            self.assertEqual(carlos["caminho"], dora["caminho"])
            wb = openpyxl.load_workbook(ana["caminho"])
            rows = list(wb["LEADS"].values)
            self.assertEqual(len(rows), 5)
            self.assertEqual([r[8] for r in rows[2:]], ["V1", "V2", "V3"])
            self.assertEqual([r[10] for r in rows[2:]], ["E1", "E2", "E3"])
            self.assertEqual(rows[2][3], "123")
            self.assertEqual(rows[4][-1], "BONIFICACAO")
            wb.close()
            enviados = []
            def enviar(*args):
                enviados.append(args)
                if args[1] == "Bia":
                    raise RuntimeError("Falha simulada")
            progresso = []
            with patch.object(gestores, "_enviar_outlook", side_effect=enviar):
                resultado = gestores.enviar_emails_gestores(selecionados, relatorios, lambda *args: progresso.append(args))
            self.assertEqual(resultado[:2], (2, 2))
            self.assertEqual(len(enviados), 3)
            self.assertEqual(enviados[0][2], ["AUTO", "VIDA"])
            self.assertEqual(progresso[-1][:2], (4, 4))

    def test_produtos_selecionados_limitam_relatorio(self):
        config = {}
        gestores.salvar_cadastro_gestor(config, self.pessoa())
        selecionados = gestores.obter_configuracoes_gestores(config).to_dict("records")[:1]
        with tempfile.TemporaryDirectory() as pasta:
            relatorios = gestores.gerar_relatorios_gerais([
                {"produto":"AUTO", "responsavel":"V1"}, {"produto":"VIDA", "responsavel":"V2"}], selecionados, pasta)
            self.assertEqual(next(iter(relatorios.values()))["quantidade"], 1)

BASELINE = 'def consolidar_remessa(\n    config: dict, fila: list[ItemFila]\n) -> tuple[list[dict], bool, list[str]]:\n    """Processa a fila de remessa: abre cada matriz, marca as linhas puxadas\n    (responsavel/modelo/datas) e monta os registros consolidados.\n\n    Retorna (registros, tem_auto_bonif, avisos).\n    """\n    registros: list[dict] = []\n    avisos: list[str] = []\n\n    tem_auto_bonif = any(item.produto.strip().upper() == "AUTO BONIF" for item in fila)\n\n    for item in fila:\n        caminho_arq = obter_caminho_matriz(config, item.produto)\n        if not caminho_arq or not os.path.isfile(caminho_arq):\n            avisos.append(\n                f"Matriz não encontrada para o produto \'{item.produto}\'. Item ignorado."\n            )\n            continue\n\n        wb = openpyxl.load_workbook(caminho_arq)  # mantém fórmulas eventuais intactas\n        ws = wb.worksheets[0]\n\n        modelo_nome = item.modelo.strip().upper()\n        qtd_coletada = 0\n\n        for row_cells in ws.iter_rows(min_row=2):\n            if qtd_coletada >= item.quantidade:\n                break\n\n            def val(col):\n                return row_cells[col - 1].value\n\n            nome = str(val(COL_MATRIZ_NOME) or "").strip()\n            tel = str(val(COL_MATRIZ_TELEFONE) or "").strip()\n            email = str(val(COL_MATRIZ_EMAIL) or "").strip()\n            if not (nome or tel or email):\n                continue\n\n            resp_existente = str(val(COL_MATRIZ_RESPONSAVEL) or "").strip()\n            dia_recolhe_existente = str(val(COL_MATRIZ_DIA_RECOLHE) or "").strip()\n            dt_entrega_existente = str(val(COL_MATRIZ_DATA_ENTREGA) or "").strip()\n            if not (\n                resp_existente == ""\n                and dia_recolhe_existente == ""\n                and dt_entrega_existente == ""\n            ):\n                continue\n\n            modelo_matriz_existente = str(val(COL_MATRIZ_MODELO) or "").strip().upper()\n            if not (\n                modelo_matriz_existente == "" or modelo_matriz_existente == modelo_nome\n            ):\n                continue\n\n            # --- marca a linha como puxada ---\n            data_entrega_dt = datetime.now().date()\n            # Weekday(vbMonday-based): segunda=1..domingo=7. Python weekday(): segunda=0..domingo=6\n            dias_add = (\n                5 if data_entrega_dt.weekday() >= 2 else 3\n            )  # >=3 no VBA (1-based) == >=2 (0-based)\n            dia_recolhe_dt = data_entrega_dt + timedelta(days=dias_add)\n\n            row_cells[COL_MATRIZ_RESPONSAVEL - 1].value = item.vendedor\n            row_cells[COL_MATRIZ_MODELO - 1].value = modelo_nome\n            row_cells[COL_MATRIZ_DATA_ENTREGA - 1].value = data_entrega_dt.strftime(\n                "%Y-%m-%d"\n            )\n            row_cells[COL_MATRIZ_DIA_RECOLHE - 1].value = dia_recolhe_dt.strftime(\n                "%Y-%m-%d"\n            )\n\n            # --- monta o registro consolidado ---\n            doc_raw = str(val(COL_MATRIZ_DOC) or "").strip()\n            doc_limpo = (\n                doc_raw.replace(".", "")\n                .replace("-", "")\n                .replace("/", "")\n                .replace(" ", "")\n            )\n\n            if len(doc_limpo) > 11:\n                colar_cpf, colar_cnpj = "", doc_raw\n            else:\n                colar_cpf, colar_cnpj = doc_raw, ""\n\n            chegada_lead = val(COL_MATRIZ_CHEGADA_LEAD)\n            chegada_lead_fmt = (\n                formatar_data_estrita(chegada_lead) if chegada_lead else ""\n            )\n\n            registro = {\n                "n": val(COL_MATRIZ_N),\n                "nome": nome,\n                "telefone": tel,\n                "colarCpf": colar_cpf,\n                "colarCnpj": colar_cnpj,\n                "email": email,\n                "chegadaDoLead": chegada_lead_fmt,\n                "diaRecolhe": dia_recolhe_dt.strftime("%Y-%m-%d"),\n                "responsavel": item.vendedor,\n                "produto": val(COL_MATRIZ_PRODUTO) or item.produto,\n                "empresa": val(COL_MATRIZ_EMPRESA) or "",\n                "modelo": modelo_nome,\n                "cpf": formatar_cpf(colar_cpf),\n                "cnpj": formatar_cnpj(colar_cnpj),\n            }\n            if tem_auto_bonif:\n                registro[HEADER_BONIF] = (\n                    "BONIFICACAO"\n                    if item.produto.strip().upper() == "AUTO BONIF"\n                    else ""\n                )\n\n            registros.append(registro)\n            qtd_coletada += 1\n\n        wb.save(caminho_arq)\n        wb.close()\n\n        if qtd_coletada < item.quantidade:\n            avisos.append(\n                f"Produto \'{item.produto}\' / modelo \'{item.modelo}\': solicitados {item.quantidade}, "\n                f"encontrados apenas {qtd_coletada} leads elegíveis."\n            )\n\n    return registros, tem_auto_bonif, avisos'

class ImportacaoTests(unittest.TestCase):
    def matriz(self, caminho, tamanho=14):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([f"C{i}" for i in range(1, 14)])
        for i in range(1, tamanho + 1):
            ws.append([i, f"Cliente {i}", "119999", "123.456.789-01", "c@test", "2026-09-30",
                       "", "", "AUTO", "E1", ["", "A", "B"][i % 3], "", "=1+1"])
        ws.cell(3, 8, "JA ENTREGUE")
        wb.save(caminho)
        wb.close()

    def test_pasta_validacao_e_exportacao(self):
        with tempfile.TemporaryDirectory() as pasta:
            destino = str(Path(pasta) / "saida nova")
            config = {"pasta_destino":destino}
            self.assertEqual(core.validar_pasta_destino(config), destino)
            primeiro = core.salvar_planilha_importacao(config, [{"nome":"Teste"}], False)
            segundo = core.salvar_planilha_importacao(config, [{"nome":"Teste"}], False)
            self.assertNotEqual(primeiro, segundo)
            self.assertTrue(Path(primeiro).is_relative_to(destino))
            with patch("tempfile.TemporaryFile", side_effect=PermissionError("negado")):
                with self.assertRaisesRegex(OSError, "acesso de escrita"):
                    core.validar_pasta_destino(config)
        with self.assertRaises(ValueError):
            core.validar_pasta_destino({"pasta_destino":""})

    def test_distribuicao_identica_ao_original_e_menos_aberturas(self):
        namespace = dict(vars(core))
        exec(BASELINE, namespace)
        original = namespace["consolidar_remessa"]
        with tempfile.TemporaryDirectory() as pasta:
            raiz = Path(pasta)
            for nome in ("antiga1", "antiga2", "nova1", "nova2"):
                self.matriz(raiz / (nome + ".xlsx"))
            def config(prefixo):
                return {"matrizes":[{"exibicao":"P1", "caminho":str(raiz / (prefixo + "1.xlsx"))},
                                    {"exibicao":"P2", "caminho":str(raiz / (prefixo + "2.xlsx"))},
                                    {"exibicao":"AUTO BONIF", "caminho":str(raiz / (prefixo + "1.xlsx"))}]}
            fila = [core.ItemFila("P1", "A", "V1", 3), core.ItemFila("P2", "B", "V2", 2),
                    core.ItemFila("P1", "B", "V3", 3), core.ItemFila("AUTO BONIF", "A", "V4", 2),
                    core.ItemFila("P2", "A", "V5", 3), core.ItemFila("P1", "B", "V6", 99)]
            with patch.object(openpyxl, "load_workbook", wraps=openpyxl.load_workbook) as leitor:
                esperado = original(config("antiga"), fila)
                self.assertEqual(leitor.call_count, 6)
            with patch.object(openpyxl, "load_workbook", wraps=openpyxl.load_workbook) as leitor:
                resultado = core.consolidar_remessa(config("nova"), fila)
                self.assertEqual(leitor.call_count, 2)
            comparacao = ([{k:v for k,v in r.items() if k != "matrizOrigem"} for r in resultado[0]], resultado[1], resultado[2])
            self.assertEqual(comparacao, esperado)
            for i in (1, 2):
                wb1 = openpyxl.load_workbook(raiz / f"antiga{i}.xlsx")
                wb2 = openpyxl.load_workbook(raiz / f"nova{i}.xlsx")
                self.assertEqual(list(wb1.active.values), list(wb2.active.values))
                wb1.close()
                wb2.close()

if __name__ == "__main__":
    unittest.main()
