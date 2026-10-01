import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
import openpyxl
import core
import envio_goalfy

class ProtecaoELeituraLeveTests(unittest.TestCase):
    def registros(self):
        return [{"n":1,"nome":"A","responsavel":"V1","produto":"AUTO"},
                {"n":2,"nome":"B","responsavel":"V2","produto":"AUTO"}]

    def test_bloqueio_persistente_e_copia_da_planilha(self):
        registros = self.registros()
        with tempfile.TemporaryDirectory() as pasta:
            historico = str(Path(pasta)/"historico.json")
            with patch.object(core.requests,"post",return_value=type("Resposta",(),{"status_code":200})()) as post:
                self.assertEqual(envio_goalfy.enviar_protegido({"webhook_url":"https://example.test"},registros,caminho_historico=historico)[:2],(2,0))
                self.assertEqual(post.call_count,2)
                with self.assertRaises(envio_goalfy.EnvioRepetidoError):
                    envio_goalfy.enviar_protegido({},copy.deepcopy(registros),caminho_historico=historico)
                self.assertEqual(post.call_count,2)
                self.assertEqual(envio_goalfy.indices_pendentes(registros,historico),[])
            cfg={"pasta_destino":pasta}
            completo=[{**{h:"" for h in core.HEADERS_BASE},**r} for r in registros]
            arquivo=core.salvar_planilha_importacao(cfg,completo,False)
            self.assertEqual(envio_goalfy.identidade_importacao(core.ler_planilha_importacao(arquivo)),envio_goalfy.identidade_importacao(completo))

    def test_falha_reenvia_somente_pendentes(self):
        registros=self.registros()
        with tempfile.TemporaryDirectory() as pasta:
            historico=str(Path(pasta)/"hist.json")
            respostas=[type("R",(),{"status_code":200})(),type("R",(),{"status_code":500})(),type("R",(),{"status_code":200})()]
            with patch.object(core.requests,"post",side_effect=respostas) as post:
                self.assertEqual(envio_goalfy.enviar_protegido({},registros,caminho_historico=historico)[:2],(1,1))
                self.assertEqual(envio_goalfy.indices_pendentes(registros,historico),[1])
                self.assertEqual(envio_goalfy.enviar_protegido({},registros,caminho_historico=historico)[:2],(1,0))
                self.assertEqual(post.call_args.kwargs["json"]["nome"],"B")
                self.assertEqual(post.call_count,3)

    def test_selecao_e_historico_invalido(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho=str(Path(pasta)/"hist.json")
            with patch.object(core.requests,"post",return_value=type("R",(),{"status_code":200})()) as post:
                envio_goalfy.enviar_protegido({},self.registros(),indices=[1],caminho_historico=caminho)
                self.assertEqual(envio_goalfy.indices_pendentes(self.registros(),caminho),[0])
                self.assertEqual(post.call_count,1)
            Path(caminho).write_text("nao e JSON",encoding="utf-8")
            with patch.object(core.requests,"post") as post:
                with self.assertRaises(ValueError):
                    envio_goalfy.enviar_protegido({},self.registros(),caminho_historico=caminho)
                self.assertEqual(post.call_count,0)

    def test_partes_excel_preservadas_e_celulas_ausentes_inseridas(self):
        with tempfile.TemporaryDirectory() as pasta:
            arquivo=Path(pasta)/"matriz.xlsx"
            wb=openpyxl.Workbook()
            ws=wb.active
            ws.append(["N","NOME","TELEFONE","CPF","EMAIL","CHEGADA","RECOLHE","RESP","PRODUTO","EMPRESA","MODELO","ENTREGA","FORMULA"])
            ws.append([1,"Cliente","11999","123",None,None,None,None,"AUTO","E1",None,None,"=1+1"])
            ws["C2"].number_format="@"
            ws["M2"].font=openpyxl.styles.Font(bold=True)
            wb.create_sheet("Outra")["A1"]="Preservar"
            wb.save(arquivo)
            wb.close()
            with ZipFile(arquivo) as z:
                antes={nome:z.read(nome) for nome in z.namelist()}
            cfg={"matrizes":[{"exibicao":"AUTO INRI","caminho":str(arquivo)}]}
            registros,_,_=core.consolidar_remessa(cfg,[core.ItemFila("AUTO INRI","LEAD","V1",1)])
            self.assertEqual(len(registros),1)
            with ZipFile(arquivo) as z:
                self.assertEqual(set(antes),set(z.namelist()))
                for nome,dados in antes.items():
                    if nome!="xl/worksheets/sheet1.xml":
                        self.assertEqual(z.read(nome),dados,nome)
            wb=openpyxl.load_workbook(arquivo)
            self.assertEqual(wb.active["H2"].value,"V1")
            self.assertEqual(wb.active["K2"].value,"LEAD")
            self.assertEqual(wb.active["M2"].value,"=1+1")
            self.assertTrue(wb.active["M2"].font.bold)
            self.assertEqual(wb["Outra"]["A1"].value,"Preservar")
            wb.close()

    def test_reenvio_manual_repetido_sem_consultar_trava(self):
        from types import SimpleNamespace
        import app
        registros=self.registros()
        estado=SimpleNamespace(after=lambda atraso,func:func(), _finalizar_reenvio_goalfy=lambda:None)
        with patch.object(core,"enviar_para_goalfy",return_value=(1,0,[])) as enviar, \
             patch.object(envio_goalfy,"enviar_protegido",side_effect=AssertionError("Reenvio manual não deve consultar a trava")), \
             patch.object(app.messagebox,"showinfo"):
            for _ in range(2):
                app.AppImportador._executar_reenvio_goalfy(estado,{},registros,indices=[1])
            self.assertEqual(enviar.call_count,2)
            self.assertEqual(enviar.call_args.args[1],[registros[1]])

    def test_layout_exato_sem_listras_ou_indice_no_titulo(self):
        import gestores
        leads=[{"nome":"Cliente","produto":"AUTO","empresa":"E1","responsavel":"V1"}]
        selecionados=[{"ID":"ana","PESSOA":"Ana","EMAIL":"ana@test","PRODUTO":"AUTO"}]
        with tempfile.TemporaryDirectory() as pasta:
            relatorio=gestores.gerar_relatorios_gerais(leads,selecionados,pasta)["ana"]
            self.assertFalse(Path(relatorio["caminho"]).stem.endswith(" - 1"))
            wb=openpyxl.load_workbook(relatorio["caminho"])
            ws=wb.active
            self.assertEqual(ws["A1"].font.color.rgb,"00000000")
            self.assertEqual(ws["A1"].fill.fgColor.rgb,"00FFFFFF")
            self.assertFalse(ws.sheet_view.showRowColHeaders)
            self.assertFalse(ws.tables["TabelaLeadsGestor"].tableStyleInfo.showRowStripes)
            self.assertIsNone(ws.tables["TabelaLeadsGestor"].tableStyleInfo.name)
            for row in ws.iter_rows(min_row=3,max_row=3):
                for cell in row:
                    self.assertEqual(cell.alignment.horizontal,"center")
                    self.assertEqual(cell.fill.fgColor.rgb,"00FFFFFF")
                    self.assertTrue(all(getattr(cell.border,lado).color.rgb=="00000000" for lado in ("top","bottom","left","right")))
            wb.close()

if __name__=="__main__":
    unittest.main()
