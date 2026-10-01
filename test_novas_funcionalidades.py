import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import openpyxl
import core
import gestores
import app

class NovasFuncionalidadesTests(unittest.TestCase):
    def test_matriz_nome_completo_sem_misturar_empresas(self):
        config = {"gestores":[{"ID":"ana", "PESSOA":"Ana", "EMAIL":"a@test", "PRODUTOS":[
            {"PRODUTO":"CELULAR INRI", "TIPO":"GESTOR", "PADRÃO":"SIM", "ATIVO":"SIM"}]}]}
        leads = [{"produto":"CELULAR", "empresa":"INRI", "matrizOrigem":"CELULAR INRI", "responsavel":"V1"},
                 {"produto":"CELULAR", "empresa":"ICX", "matrizOrigem":"CELULAR ICX", "responsavel":"V2"}]
        df, ausentes = gestores.destinatarios_por_produto(leads, gestores.obter_configuracoes_gestores(config))
        self.assertEqual(list(df["PRODUTO"]), ["CELULAR INRI"])
        self.assertEqual(ausentes, ["CELULAR ICX"])
        with tempfile.TemporaryDirectory() as pasta:
            relatorios = gestores.gerar_relatorios_gerais(leads, df.to_dict("records"), pasta)
            self.assertEqual(relatorios["ana"]["quantidade"], 1)
            wb = openpyxl.load_workbook(relatorios["ana"]["caminho"])
            self.assertEqual(wb.active["K3"].value, "INRI")
            self.assertEqual(len(wb.active.tables), 1)
            self.assertTrue(wb.active["A1"].value.startswith("RELATÓRIO GERAL - ENTREGA DE LEADS - "))
            for row in wb.active.iter_rows(min_row=2, max_row=3):
                for cell in row:
                    self.assertTrue(all(getattr(cell.border, lado).style == "thin" for lado in ("top","bottom","left","right")))
            wb.close()

    def test_tabela_e_bordas_titulo_vendedor(self):
        with tempfile.TemporaryDirectory() as pasta:
            arquivo = core.gerar_planilha_grupo(("V1", "AUTO", "INRI"), [{"nome":"Cliente", "responsavel":"V1", "produto":"AUTO"}], pasta)
            wb = openpyxl.load_workbook(arquivo)
            ws = wb.active
            self.assertEqual(len(ws.tables), 1)
            self.assertEqual(ws.tables["TabelaLeadsVendedor"].ref, "B3:M4")
            self.assertEqual(ws["B2"].border.left.style, "thick")
            self.assertEqual(ws["M2"].border.right.style, "thick")
            for col in range(2,14):
                self.assertEqual(ws.cell(2,col).border.top.style, "thick")
                self.assertEqual(ws.cell(2,col).border.bottom.style, "thick")
            for row in ws.iter_rows(min_row=3,max_row=4,min_col=2,max_col=13):
                for cell in row:
                    self.assertTrue(all(getattr(cell.border, lado).style == "thin" for lado in ("top","bottom","left","right")))
            wb.close()

    def test_leitura_importacao_salva_preserva_payload_e_origem(self):
        with tempfile.TemporaryDirectory() as pasta:
            registro = {h:"" for h in core.HEADERS_BASE}
            registro.update(nome="Cliente", telefone="00119999", cpf="00123456789", produto="CELULAR", empresa="INRI", responsavel="V1", bonif="BONIFICACAO", matrizOrigem="CELULAR INRI")
            arquivo = core.salvar_planilha_importacao({"pasta_destino":pasta}, [registro], True)
            antes = Path(arquivo).read_bytes()
            lidos = core.ler_planilha_importacao(arquivo)
            self.assertEqual(core.montar_payload(lidos[0]), core.montar_payload(registro))
            self.assertEqual(lidos[0]["matrizOrigem"], "CELULAR INRI")
            self.assertEqual(lidos[0]["_linha_origem"], 2)
            self.assertEqual(Path(arquivo).read_bytes(), antes)
            errado = Path(pasta)/"errado.xlsx"
            wb = openpyxl.Workbook()
            wb.save(errado)
            wb.close()
            with self.assertRaises(ValueError):
                core.ler_planilha_importacao(errado)

    def test_segunda_geracao_bloqueada_antes_de_alterar_matrizes(self):
        fila = [core.ItemFila("AUTO", "A", "V1", 3)]
        assinatura = (("AUTO", "A", "V1", 3),)
        estado = SimpleNamespace(_gerando=False, fila=fila, _exportacao_pendente=None,
                                 travar_repeticao=SimpleNamespace(get=lambda:True), _assinaturas_geradas={assinatura})
        with patch.object(app.messagebox,"showwarning") as aviso, patch.object(core,"consolidar_remessa") as consolidar:
            app.AppImportador._gerar_importacao_thread(estado)
            self.assertEqual(aviso.call_count, 1)
            self.assertEqual(consolidar.call_count, 0)

if __name__ == "__main__":
    unittest.main()
