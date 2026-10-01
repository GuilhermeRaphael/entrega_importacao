"""Leitura leve e alteração pontual da primeira aba, sem reescrever a matriz inteira."""
import os
import re
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet._reader import WorkSheetParser

CELULA = re.compile(rb'<c\b(?P<attrs>[^>]*?)(?:/>|>(?P<body>.*?)</c>)', re.S)
LINHA = re.compile(rb'(?P<abertura><row\b[^>]*\br="(?P<numero>\d+)"[^>]*>)(?P<conteudo>.*?)</row>', re.S)
REFERENCIA = re.compile(rb'\br="([A-Z]+)(\d+)"')

class CelulaLeve:
    def __init__(self, matriz, row, column, value):
        self.matriz, self.row, self.column, self._value = matriz, row, column, value
    @property
    def value(self):
        return self._value
    @value.setter
    def value(self, value):
        if value != self._value:
            self.matriz.alteracoes.setdefault(self.row, {})[self.column] = value
            self._value = value

class MatrizLeve:
    def __init__(self, caminho, elegivel):
        self.caminho = caminho
        stat = os.stat(caminho)
        self.estado_original = (stat.st_size, stat.st_mtime_ns)
        self.alteracoes = {}
        self._cells = {}
        self.worksheets = [self]
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=False)
        try:
            ws = wb.worksheets[0]
            self.aba_xml = ws._worksheet_path.lstrip('/')
            with ws._get_source() as fonte:
                parser = WorkSheetParser(fonte, ws._shared_strings, data_only=False,
                                         epoch=wb.epoch, date_formats=wb._date_formats,
                                         timedelta_formats=wb._timedelta_formats)
                for row, cells in parser.parse():
                    if row < 2:
                        continue
                    valores = {c['column']:c['value'] for c in cells if c['column'] <= 12}
                    linha = tuple(valores.get(i) for i in range(1,13))
                    if not elegivel(linha):
                        continue
                    for i, value in enumerate(linha, 1):
                        self._cells[(row,i)] = CelulaLeve(self,row,i,value)
        finally:
            wb.close()

    def cell(self,row,column):
        chave = (row,column)
        if chave not in self._cells:
            self._cells[chave] = CelulaLeve(self,row,column,None)
        return self._cells[chave]

    def close(self):
        pass

    def salvar(self, caminho):
        if not self.alteracoes:
            return
        if os.path.normcase(os.path.abspath(caminho)) != os.path.normcase(os.path.abspath(self.caminho)):
            raise ValueError('A atualização da matriz deve ser feita no arquivo de origem.')
        stat = os.stat(caminho)
        if (stat.st_size,stat.st_mtime_ns) != self.estado_original:
            raise RuntimeError('A matriz foi alterada por outro processo durante a entrega. Feche o Excel e tente novamente.')
        encontradas = set()
        def montar(coluna, numero, value, attrs=None):
            if not isinstance(value,str):
                raise TypeError('A marcação da entrega deve ser textual.')
            ref = f'{get_column_letter(coluna)}{numero}'.encode()
            if attrs is None:
                attrs = b' r="' + ref + b'"'
            attrs = re.sub(rb'\s+t="[^"]*"', b'', attrs)
            return b'<c' + attrs + b' t="inlineStr"><is><t xml:space="preserve">' + escape(value).encode('utf-8') + b'</t></is></c>'
        def alterar_linha(match):
            numero = int(match['numero'])
            if numero not in self.alteracoes:
                return match.group(0)
            encontradas.add(numero)
            pendentes = dict(self.alteracoes[numero])
            conteudo = match['conteudo']
            partes, cursor = [], 0
            for cell in CELULA.finditer(conteudo):
                ref = REFERENCIA.search(cell['attrs'])
                if ref is None:
                    continue
                coluna = column_index_from_string(ref[1].decode())
                partes.append(conteudo[cursor:cell.start()])
                for menor in sorted(k for k in pendentes if k < coluna):
                    partes.append(montar(menor,numero,pendentes.pop(menor)))
                if coluna in pendentes:
                    partes.append(montar(coluna,numero,pendentes.pop(coluna),cell['attrs']))
                else:
                    partes.append(cell.group(0))
                cursor = cell.end()
            for coluna in sorted(pendentes):
                partes.append(montar(coluna,numero,pendentes[coluna]))
            partes.append(conteudo[cursor:])
            return match['abertura'] + b''.join(partes) + b'</row>'
        temp = None
        try:
            with ZipFile(caminho) as origem:
                xml_original = origem.read(self.aba_xml)
                xml = LINHA.sub(alterar_linha,xml_original)
                if encontradas != set(self.alteracoes):
                    raise ValueError('Formato XML da matriz não suportado para atualização pontual.')
                from xml.etree import ElementTree
                ElementTree.fromstring(xml)
                with tempfile.NamedTemporaryFile(dir=os.path.dirname(caminho),prefix='.entrega_',suffix='.xlsx',delete=False) as arq:
                    temp = arq.name
                with ZipFile(temp,'w',compression=ZIP_DEFLATED,compresslevel=1,allowZip64=True) as destino:
                    destino.comment = origem.comment
                    for info in origem.infolist():
                        dados = xml if info.filename == self.aba_xml else origem.read(info.filename)
                        destino.writestr(info,dados,compress_type=info.compress_type,compresslevel=1)
            stat = os.stat(caminho)
            if (stat.st_size,stat.st_mtime_ns) != self.estado_original:
                raise RuntimeError('A matriz foi alterada durante o salvamento. Nenhuma atualização foi aplicada.')
            os.replace(temp,caminho)
            temp = None
        finally:
            if temp and os.path.exists(temp):
                os.remove(temp)
