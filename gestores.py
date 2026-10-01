"""Seleção e envio de relatórios para gestores, independente da distribuição."""
from pathlib import Path
import logging
import tempfile
import unicodedata
import pandas as pd
import core

COLUNAS = ["PRODUTO", "PESSOA", "TIPO", "PADRÃO", "ATIVO", "EMAIL"]
logger = logging.getLogger(__name__)


def _texto(valor):
    return "" if valor is None or pd.isna(valor) else str(valor).strip()


def _chave(valor):
    return "".join(c for c in unicodedata.normalize("NFD", _texto(valor).upper())
                   if unicodedata.category(c) != "Mn")


def migrar_cadastros_gestores(config):
    """Agrupa cadastros antigos por nome/e-mail sem descartar seus vínculos."""
    import uuid
    registros, agrupados = [], {}
    for antigo in config.get("gestores", []):
        if "PRODUTOS" in antigo:
            novo = dict(antigo)
            novo.setdefault("ID", uuid.uuid4().hex)
            registros.append(novo)
            continue
        chave = (_chave(antigo.get("PESSOA")), _texto(antigo.get("EMAIL")).casefold())
        if chave not in agrupados:
            agrupados[chave] = {"ID": uuid.uuid4().hex, "PESSOA": _texto(antigo.get("PESSOA")),
                                "EMAIL": _texto(antigo.get("EMAIL")), "PRODUTOS": []}
            registros.append(agrupados[chave])
        agrupados[chave]["PRODUTOS"].append({c: _texto(antigo.get(c)) for c in ("PRODUTO", "TIPO", "PADRÃO", "ATIVO")})
    config["gestores"] = registros
    return config


def obter_configuracoes_gestores(config):
    """Expande os vínculos para a seleção; identidade/e-mail ficam num único cadastro."""
    linhas = []
    for pessoa in config.get("gestores", []):
        vinculos = pessoa.get("PRODUTOS", [pessoa])
        for vinculo in vinculos:
            r = {c: _texto(vinculo.get(c, pessoa.get(c, ""))) for c in COLUNAS}
            r["ID"] = pessoa.get("ID", _chave(pessoa.get("PESSOA")))
            linhas.append(r)
    return pd.DataFrame(linhas, columns=COLUNAS + ["ID"])


def salvar_cadastro_gestor(config, cadastro, indice=None):
    import uuid
    import copy
    nome = _texto(cadastro.get("PESSOA"))
    if not nome:
        raise ValueError("Preencha o nome.")
    registro = {"ID": cadastro.get("ID") or uuid.uuid4().hex, "PESSOA": nome,
                "EMAIL": _texto(cadastro.get("EMAIL")), "PRODUTOS": copy.deepcopy(cadastro.get("PRODUTOS", []))}
    vistos = set()
    for r in registro["PRODUTOS"]:
        r["PRODUTO"] = _texto(r.get("PRODUTO"))
        chave = _chave(r["PRODUTO"])
        if not chave or chave in vistos:
            raise ValueError("Cada produto deve ter um único vínculo por gestor.")
        vistos.add(chave)
        if r.get("TIPO") not in ("GESTOR", "SUBGESTOR"):
            raise ValueError("Selecione GESTOR ou SUBGESTOR.")
        if any(r.get(c) not in ("SIM", "NÃO") for c in ("PADRÃO", "ATIVO")):
            raise ValueError("Padrão e ativo devem ser SIM ou NÃO.")
    registros = config.setdefault("gestores", [])
    if indice is not None and not 0 <= indice < len(registros):
        raise ValueError("Cadastro não encontrado.")
    for i, existente in enumerate(registros):
        if i != indice and _chave(existente.get("PESSOA")) == _chave(nome):
            raise ValueError("Essa pessoa já está cadastrada. Edite seus produtos no cadastro existente.")
    if indice is None:
        registros.append(registro)
    else:
        registro["ID"] = registros[indice]["ID"]
        registros[indice] = registro


def excluir_cadastro_gestor(config, indice):
    del config["gestores"][indice]


def produtos_do_lead(r):
    return {_chave(v) for v in (r.get("produto"), r.get("matrizOrigem")) if _texto(v)}


def destinatarios_por_produto(produtos_leads, gestores_df):
    registros = [p if isinstance(p, dict) else {"produto": p} for p in produtos_leads]
    disponiveis = set().union(*(produtos_do_lead(r) for r in registros)) if registros else set()
    ativos = gestores_df[gestores_df["ATIVO"].map(_chave).eq("SIM")]
    candidatos = ativos[ativos["PRODUTO"].map(_chave).isin(disponiveis)].copy()
    cadastrados = set(candidatos["PRODUTO"].map(_chave))
    ausentes = list(dict.fromkeys(_texto(r.get("matrizOrigem") or r.get("produto")) for r in registros
                                  if not produtos_do_lead(r).intersection(cadastrados)))
    return candidatos.reset_index(drop=True), ausentes


def abrir_interface_selecao_gestores(produtos_leads, gestores_df, parent=None):
    import customtkinter as ctk
    root = None
    if parent is None:
        root = ctk.CTk()
        root.withdraw()
        parent = root
    popup = ctk.CTkToplevel(parent)
    popup.title("Selecionar Gestores e Subgestores")
    popup.geometry("760x560")
    popup.transient(parent)
    popup.grab_set()
    resultado = None
    candidatos, ausentes = destinatarios_por_produto(produtos_leads, gestores_df)
    ctk.CTkLabel(popup, text="Entrega de leads para gestores", font=ctk.CTkFont(size=20, weight="bold")).pack(padx=20, pady=(18, 4), anchor="w")
    ctk.CTkLabel(popup, text="Selecione os destinatários e produtos. Cada pessoa receberá um relatório consolidado.", wraplength=700).pack(padx=20, pady=(0, 10), anchor="w")
    scroll = ctk.CTkScrollableFrame(popup)
    scroll.pack(padx=20, pady=5, fill="both", expand=True)
    variaveis = []
    registros = candidatos.to_dict("records")
    por_produto = {}
    for i, r in enumerate(registros):
        por_produto.setdefault(r["PRODUTO"], []).append((i, r))
    for produto, pessoas in por_produto.items():
        ctk.CTkLabel(scroll, text=produto, font=ctk.CTkFont(size=15, weight="bold"), text_color="#0066CC").pack(anchor="w", padx=10, pady=(12, 4))
        for i, r in pessoas:
            var = ctk.BooleanVar(value=_chave(r["PADRÃO"]) == "SIM")
            texto = f"{r['PESSOA']}  •  {r['TIPO']}"
            if _chave(r["PADRÃO"]) == "SIM":
                texto += "  •  Padrão"
            ctk.CTkCheckBox(scroll, text=texto, variable=var).pack(anchor="w", padx=12, pady=(5, 1))
            ctk.CTkLabel(scroll, text=r["EMAIL"] or "SEM E-MAIL CADASTRADO", text_color="gray").pack(anchor="w", padx=42, pady=(0, 5))
            variaveis.append((i, var))
    for produto in ausentes:
        ctk.CTkLabel(scroll, text=f"{produto} — (SEM GESTOR CADASTRADO)", text_color="gray").pack(anchor="w", padx=10, pady=8)
    botoes = ctk.CTkFrame(popup, fg_color="transparent")
    botoes.pack(fill="x", padx=20, pady=10)
    def marcar(valor):
        for _, var in variaveis:
            var.set(valor)
    ctk.CTkButton(botoes, text="Selecionar Todos", command=lambda: marcar(True), fg_color="#003366").pack(side="left", padx=(0, 8))
    ctk.CTkButton(botoes, text="Limpar Seleção", command=lambda: marcar(False), fg_color="#555555").pack(side="left")
    def confirmar():
        nonlocal resultado
        resultado = [registros[i] for i, var in variaveis if var.get()]
        popup.destroy()
    rodape = ctk.CTkFrame(popup, fg_color="transparent")
    rodape.pack(fill="x", padx=20, pady=(0, 18))
    ctk.CTkButton(rodape, text="CONFIRMAR ENVIO", height=38, fg_color="#0066CC", command=confirmar).pack(side="left", fill="x", expand=True, padx=(0, 8))
    ctk.CTkButton(rodape, text="CANCELAR", height=38, fg_color="#B42828", command=popup.destroy).pack(side="right")
    parent.wait_window(popup)
    if root:
        root.destroy()
    return resultado


def gerar_relatorios_gerais(leads, selecionados, pasta_saida):
    """Um relatório consolidado por conjunto de produtos, reutilizado entre pessoas."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    if isinstance(leads, pd.DataFrame):
        leads = leads.to_dict("records")
    pessoas = {}
    for r in selecionados:
        chave = r.get("ID") or (_chave(r["PESSOA"]), _texto(r.get("EMAIL")).casefold())
        pessoa = pessoas.setdefault(chave, {"cadastro": r, "produtos": set()})
        pessoa["produtos"].add(_chave(r["PRODUTO"]))
    cache, relatorios = {}, {}
    campos = ["n", "nome", "telefone", "cpf", "cnpj", "email", "chegadaDoLead", "diaRecolhe",
              "responsavel", "produto", "empresa", "modelo"]
    cabecalhos = list(core.HEADERS_GRUPO)
    if any("bonif" in r for r in leads):
        campos.append("bonif")
        cabecalhos.append("BONIFICAÇÃO")
    for chave, pessoa in pessoas.items():
        conjunto = tuple(sorted(pessoa["produtos"]))
        rows = [r for r in leads if produtos_do_lead(r).intersection(pessoa["produtos"])]
        if not rows:
            continue
        if conjunto not in cache:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "LEADS"
            from datetime import datetime
            from openpyxl.styles import Alignment, Border, Side
            ultima_coluna = get_column_letter(len(campos))
            ws.merge_cells(f"A1:{ultima_coluna}1")
            ws["A1"] = f"RELATÓRIO GERAL - ENTREGA DE LEADS - {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="000000")
            ws["A1"].fill = PatternFill("solid", fgColor="FFFFFF")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            grosso = Side(style="thick", color="000000")
            for i in range(1, len(campos) + 1):
                ws.cell(1, i).border = Border(top=grosso, bottom=grosso, left=grosso if i == 1 else Side(), right=grosso if i == len(campos) else Side())
            ws.row_dimensions[1].height = 32
            ws.append(cabecalhos)
            larguras = [len(h) for h in cabecalhos]
            for r in rows:
                valores = [r.get(c, "") for c in campos]
                ws.append(valores)
                larguras = [max(a, len(str(v or ""))) for a, v in zip(larguras, valores)]
            for cell in ws[2]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="002060")
            for i, largura in enumerate(larguras, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(largura + 2, 40)
            core.formatar_tabela_relatorio(ws, 2, ws.max_row, 1, len(campos), "TabelaLeadsGestor")
            Path(pasta_saida).mkdir(parents=True, exist_ok=True)
            pasta_relatorio = Path(pasta_saida) / f"anexo_{len(cache)}"
            pasta_relatorio.mkdir(parents=True, exist_ok=True)
            caminho = str(pasta_relatorio / f"RELATÓRIO GERAL - ENTREGA DE LEADS - {datetime.now().strftime('%d-%m-%Y %Hh%M')}.xlsx")
            core.salvar_workbook_rapido(wb, caminho)
            wb.close()
            cache[conjunto] = caminho
        relatorios[chave] = {"cadastro": pessoa["cadastro"], "caminho": cache[conjunto],
                            "produtos": sorted({_texto(r.get("produto")) for r in rows}), "quantidade": len(rows)}
    return relatorios


def _enviar_outlook(destinatario, nome, produtos, caminho):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        mail = outlook.CreateItem(0)
        mail.To = destinatario
        mail.Subject = "Relatório geral de entrega de leads"
        mail.Body = (f"Olá {nome},\n\nSegue em anexo o relatório geral dos leads entregues "
                     f"referentes a {', '.join(produtos)}, com os clientes, vendedores, produtos, "
                     "empresas e demais informações da entrega.")
        mail.Attachments.Add(str(Path(caminho).resolve()))
        mail.Send()
    finally:
        pythoncom.CoUninitialize()


def enviar_emails_gestores(gestores_selecionados, relatorios_gerados, progress_callback=None):
    """Um e-mail e um anexo consolidado por pessoa; falhas não interrompem os demais."""
    sucesso, falha, avisos = 0, 0, []
    for i, relatorio in enumerate(relatorios_gerados.values(), 1):
        r = relatorio["cadastro"]
        ok = False
        try:
            email = _texto(r.get("EMAIL"))
            if not email:
                raise ValueError("sem e-mail cadastrado; envio ignorado")
            _enviar_outlook(email, r["PESSOA"], relatorio["produtos"], relatorio["caminho"])
            sucesso += 1
            ok = True
        except Exception as exc:
            falha += 1
            aviso = f"{r['PESSOA']}: {exc}"
            avisos.append(aviso)
            logger.warning(aviso)
        if progress_callback:
            progress_callback(i, len(relatorios_gerados), ok)
    return sucesso, falha, avisos
